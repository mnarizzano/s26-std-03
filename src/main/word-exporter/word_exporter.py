
import os
import json
from docx import Document
from docx.shared import Inches

try:
    from disassembly_loader import build_guide
except ImportError:
    try:
        from loader_se import build_guide
    except ImportError:
        build_guide = None


def _extract_name(entity):
    # extracts a readable string name from object / tuple / dictionary
    if not entity:
        return ""
    if isinstance(entity, (tuple, list)):
        if len(entity) == 0:
            return ""
        entity = entity[0]
    if hasattr(entity, "name"):
        return str(entity.name)
    if isinstance(entity, dict) and "name" in entity:
        return str(entity["name"])
    return str(entity)

# exports the guide in docx (see README for the explanation of the parameters)
def export_to_word(json_path, output_path=None, depth=None, include_bom=False, **kwargs):
    
    if output_path is not None and not isinstance(output_path, (str, bytes, os.PathLike)):
        include_bom = depth if depth is not None else include_bom
        depth = output_path
        output_path = None

    if not build_guide:
        raise ImportError("Could not import 'build_guide' from disassembly loader.")

    guide = build_guide(json_path, depth=depth, include_bom=include_bom)
    
    # a map to read the json and to understand the component node of each action (without always coming back to the root component)
    operation_to_source = {}
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            j_data = json.load(f)
            shapes = {s['id']: s for s in j_data.get('shapes', [])}
            for s in j_data.get('shapes', []):
                if s.get('type') == 'arrow':
                    from_shape = shapes.get(s.get('from_shape_id'))
                    to_shape = shapes.get(s.get('to_shape_id'))
                    # if i go from a component (rectangle) to an action (diamond) i save the couple
                    if from_shape and to_shape and from_shape.get('type') == 'component' and to_shape.get('type') == 'diamond':
                        op_text = to_shape.get('text', '').strip()
                        comp_name = from_shape.get('name') or from_shape.get('text', '')
                        if op_text and comp_name:
                            operation_to_source[op_text] = comp_name
    except Exception:
        pass

    doc = Document()
    doc.add_heading(f"Disassembly guide: {guide.product.name}", 0)

    # validation warnings
    if guide.warnings:
        doc.add_heading("Validation warnings", level=1)
        for warning in guide.warnings:
            p_warn = doc.add_paragraph(style='Normal')
            r_warn = p_warn.add_run(f"[{warning.severity.upper()}] Rule '{warning.rule}': {warning.message}")
            r_warn.italic = True

    # bill of materials
    if hasattr(guide, 'bill_of_materials') and guide.bill_of_materials:
        doc.add_heading("Bill of materials", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Component'
        hdr_cells[1].text = 'Weight'
        hdr_cells[2].text = 'Material'
        hdr_cells[3].text = 'Color'

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for item in guide.bill_of_materials:
            row_cells = table.add_row().cells
            row_cells[0].text = _extract_name(item)

            if hasattr(item, 'weight') and item.weight is not None:
                val = item.weight
                display_weight = int(val) if isinstance(val, float) and val.is_integer() else val
                row_cells[1].text = f"{display_weight} g"
            else:
                row_cells[1].text = "—"

            row_cells[2].text = str(item.material) if hasattr(item, 'material') and item.material is not None else "—"
            row_cells[3].text = str(item.color) if hasattr(item, 'color') and item.color is not None else "—"

        # fixed width columns
        col_widths = [Inches(3.0), Inches(1.0), Inches(1.0), Inches(1.0)]
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width
        doc.add_paragraph()

    # i save the active subroots to know what i have to tell in the "Then resume from..."
    active_subroots = set()
    for i, step in enumerate(guide.steps):
        op_text = getattr(step, 'operation', '').strip()
        source_name = operation_to_source.get(op_text)
        
        # fallback if not found in the input Json
        if not source_name:
            if i == 0 or not getattr(guide.steps[i - 1], 'continues_as', None):
                source_name = guide.product.name
            else:
                source_name = _extract_name(guide.steps[i - 1].continues_as)
                
        if source_name:
            active_subroots.add(source_name)

    # steps
    for i, step in enumerate(guide.steps):
        op_text = getattr(step, 'operation', '').strip()
        current_source = operation_to_source.get(op_text)
        
        # another fallback 
        if not current_source:
            if i == 0 or not getattr(guide.steps[i - 1], 'continues_as', None):
                current_source = guide.product.name
            else:
                current_source = _extract_name(guide.steps[i - 1].continues_as)

        doc.add_paragraph()
        doc.add_heading(f"Disassembling from {current_source}", level=1)

        p_action = doc.add_paragraph()
        r_action = p_action.add_run(f"Action {step.index}: {step.operation}")
        r_action.bold = True

        if getattr(step, 'tools_required', None):
            p_tools = doc.add_paragraph()
            r_tools_label = p_tools.add_run("Tools required: ")
            r_tools_label.bold = True
            tools_clean = [_extract_name(t) for t in step.tools_required if _extract_name(t)]
            p_tools.add_run(", ".join(tools_clean))

        if getattr(step, 'actions', None):
            for action in step.actions:
                action_text = getattr(action, 'text', str(action))
                try:
                    doc.add_paragraph(action_text, style='List Bullet')
                except Exception:
                    doc.add_paragraph(f"• {action_text}")

        # i check whether the disassembled pieces will have to be disassembled as well
        collected_sub_roots = []
        continues_name = _extract_name(getattr(step, 'continues_as', None))
        if continues_name and continues_name in active_subroots:
            collected_sub_roots.append(continues_name)

        if hasattr(step, 'outputs') and step.outputs:
            for out_component in step.outputs:
                comp_name = _extract_name(out_component)
                if comp_name and comp_name in active_subroots and comp_name not in collected_sub_roots:
                    collected_sub_roots.append(comp_name)

        # print the message to tell where to resume the dismounting process
        if collected_sub_roots:
            p = doc.add_paragraph()
            resume_phrases = [f'"Disassembling from {name}"' for name in collected_sub_roots]
            r = p.add_run("Then resume from " + ", ".join(resume_phrases))
            r.bold = True

    # save file in specified path or in json path
    if not output_path:
        base_name = os.path.splitext(os.path.basename(json_path))[0]
        output_path = os.path.join(os.path.dirname(os.path.abspath(json_path)), f"{base_name}_guide.docx")

    target_dir = os.path.dirname(os.path.abspath(output_path))
    if target_dir:
        os.makedirs(target_dir, exist_ok=True)

    doc.save(output_path)
    return os.path.abspath(output_path)
